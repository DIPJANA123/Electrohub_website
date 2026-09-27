import os

from django.core.management.base import BaseCommand
from django.core.files import File
from openpyxl import load_workbook

from shop.models import Product, Category, ProductImage


class Command(BaseCommand):
    help = "Import products from Excel"

    def handle(self, *args, **kwargs):

        excel_path = os.path.join(
            os.path.dirname(
                os.path.dirname(
                    os.path.dirname(
                        os.path.dirname(__file__)
                    )
                )
            ),
            "electrohub_products.xlsx"
        )

        wb = load_workbook(excel_path, data_only=True)
        ws = wb.active

        added = 0
        skipped = 0
        images_added = 0

        for row in ws.iter_rows(min_row=2, values_only=True):

            product_number = row[0]
            name = row[1]
            category_name = row[2]
            description = row[3]
            price = row[4]
            main_image = row[5]

            if not name:
                continue

            category, created = Category.objects.get_or_create(
                name=category_name
            )

            product = Product.objects.filter(name=name).first()

            if product:
                skipped += 1
            else:
                product = Product.objects.create(
                    name=name,
                    category=category,
                    description=description or "",
                    price=price or 0,
                    is_approved=True
                )

                added += 1

                self.stdout.write(
                    f"Added: {product_number} - {name}"
                )

            # Main image
            if main_image:
                image_path = os.path.join(
                    "media",
                    "products",
                    str(main_image)
                )

                if os.path.exists(image_path):

                    if not product.image:
                        with open(image_path, "rb") as image_file:
                            product.image.save(
                                str(main_image),
                                File(image_file),
                                save=True
                            )

                        images_added += 1

                        self.stdout.write(
                            f"Image added: {name} -> {main_image}"
                        )

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS(
                f"Import complete! "
                f"Added: {added}, "
                f"Skipped: {skipped}, "
                f"Images added: {images_added}"
            )
        )